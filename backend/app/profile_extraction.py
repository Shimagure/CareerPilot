"""Conservative, deterministic resume text extraction for profile drafting.

Extraction produces *candidate* information only. It never:
- assigns a proficiency level or years of experience to a skill,
- fabricates dates, employers, titles, or metrics,
- turns a heading or generic word into a qualification,
- upgrades anything to verified status (every item is tagged
  ``resume_extracted`` until the user confirms it).

Ambiguous content is left unresolved rather than guessed.
"""
from __future__ import annotations

import io
import re
from typing import Any

try:  # pragma: no cover - import guard for environments without the dependency
    import docx  # python-docx
except Exception:  # pragma: no cover
    docx = None

try:  # pragma: no cover
    from pypdf import PdfReader
except Exception:  # pragma: no cover
    PdfReader = None

SOURCE = "resume_extracted"

_EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")
_PHONE_RE = re.compile(r"\+?[\d\s.()-]{7,24}")
_LINKEDIN_RE = re.compile(
    r"(?:https?://)?(?:www\.)?linkedin\.com/(?:in|company)/[A-Za-z0-9\-_%]+",
    re.IGNORECASE,
)
_URL_RE = re.compile(r"https?://[^\s]+", re.IGNORECASE)
_GITHUB_RE = re.compile(r"https?://(?:www\.)?github\.com/[^\s]+", re.IGNORECASE)

_MONTHS = (
    "January|February|March|April|May|June|July|August|September|October|"
    "November|December|Jan|Feb|Mar|Apr|Jun|Jul|Aug|Sep|Oct|Nov|Dec"
)
_FLEX_YEAR = rf"(?:{_MONTHS})\.?\s+(?:19|20)\d{{2}}|(?:19|20)\d{{2}}"
_RANGE_RE = re.compile(
    rf"((?:{_MONTHS})\.?\s+)?((?:19|20)\d{{2}})"
    rf"\s*(?:[-–—/]|\bto\b)\s*"
    rf"(((?:{_MONTHS})\.?\s+)?((?:19|20)\d{{2}}|present|current|now|today))",
    re.IGNORECASE,
)

_NAME_RE = re.compile(
    r"^[A-Z][A-Za-zÀ-ÿ.'-]*(?:\s+[A-Z][A-Za-zÀ-ÿ.'&-]*){1,5}$"
)
_ALL_CAPS_NAME_RE = re.compile(r"^[A-Z][A-Z\s.'&-]{2,45}$")

_DOCUMENT_HEADERS = {
    "curriculum vitae", "curriculum vitae (cv)", "cv", "résumé", "resume",
    "vitae", "personal details", "contact details", "contact information",
}

_HEADER_MAP: dict[str, tuple[str, ...]] = {
    "summary": (
        "summary", "professional summary", "profile summary", "profile",
        "objective", "career objective", "about me", "about",
    ),
    "skills": (
        "skills", "technical skills", "core skills", "key skills",
        "core competencies", "technical competencies", "technologies",
        "tools and technologies", "expertise", "areas of expertise",
    ),
    "experience": (
        "experience", "work experience", "professional experience",
        "employment", "employment history", "work history", "career history",
        "professional history", "recent experience",
    ),
    "education": (
        "education", "educational background", "academic background",
        "academic", "academics", "education and training",
    ),
    "certifications": (
        "certifications", "certification", "certificates", "certificate",
        "licenses", "licences", "professional certifications",
        "training and certifications", "trainings",
    ),
    "languages": ("languages", "language"),
    "projects": ("projects", "key projects", "personal projects"),
}

_DEGREE_KEYWORDS = (
    "bachelor", "b.s.", "bs", "bsc", "master", "m.s.", "ms", "mba", "doctor",
    "phd", "associate", "diploma", "degree", "postgraduate", "graduate",
)

_BULLET = re.compile(r"^[\s>#]*[-•*·]\s*")
_NUMBERED = re.compile(r"^\d+[.)]\s+")


def _normalize_heading(line: str) -> str:
    cleaned = line.strip().strip("*•·-–—#").strip().rstrip(":")
    cleaned = re.sub(r"^\s*\d+[.)]\s*", "", cleaned).strip()
    return cleaned.lower()


def _find_heading(line: str) -> str | None:
    cleaned = _normalize_heading(line)
    if not cleaned or len(cleaned) > 24:
        return None
    for key, tokens in _HEADER_MAP.items():
        if cleaned in tokens:
            return key
    return None


def extract_text(content: bytes, kind: str) -> tuple[str, str | None]:
    """Return (extracted_text, warning_or_None). Never raises."""
    try:
        if kind == "pdf":
            return _extract_pdf_text(content)
        return _extract_docx_text(content)
    except Exception:
        return "", "The document could not be read as text."


def _extract_pdf_text(content: bytes) -> tuple[str, str | None]:
    if PdfReader is None:
        return "", "PDF text extraction is unavailable."
    try:
        reader = PdfReader(io.BytesIO(content))
    except Exception:
        return "", "The PDF could not be read."
    pages: list[str] = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            continue
    text = "\n".join(pages)
    if not text.strip():
        return "", (
            "This PDF contains no extractable text (it may be a scanned image). "
            "You can add your profile information manually."
        )
    return text, None


def _extract_docx_text(content: bytes) -> tuple[str, str | None]:
    if docx is None:
        return "", "DOCX text extraction is unavailable."
    document = docx.Document(io.BytesIO(content))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    text = "\n".join(parts)
    if not text.strip():
        return "", "This document contains no readable text."
    return text, None


def build_profile_draft(content: bytes, kind: str) -> tuple[dict[str, Any], str | None]:
    """Extract a reviewable profile draft and an optional warning.

    Returns (draft, warning). When the document yields no usable text the draft
    is empty and ``warning`` explains why (scanned PDF etc.).
    """
    text, warning = extract_text(content, kind)
    if not text.strip():
        return _empty_draft(), warning if warning is not None else (
            "No readable text could be extracted from this document."
        )
    try:
        return _parse_text(text), warning
    except Exception:
        return _empty_draft(), (
            "The resume text could not be interpreted. You can add your profile "
            "information manually."
        )


def _empty_draft() -> dict[str, Any]:
    return {
        "full_name": None,
        "headline": None,
        "location": None,
        "phone": None,
        "email": None,
        "linkedin_url": None,
        "portfolio_url": None,
        "professional_summary": None,
        "skills": [],
        "experience": [],
        "education": [],
        "certifications": [],
        "languages": [],
        "projects": [],
    }


def _segment_sections(text: str) -> tuple[list[str], dict[str, list[str]]]:
    """Split text into a top block and ordered heading -> lines sections."""
    raw_lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    top: list[str] = []
    sections: dict[str, list[str]] = {}
    order: list[str] = []
    current: str | None = None
    for raw in raw_lines:
        line = raw.strip()
        if not line:
            if current is not None:
                sections[current].append("")
            continue
        heading = _find_heading(line)
        if heading is not None:
            current = heading
            if heading not in sections:
                sections[heading] = []
                order.append(heading)
            continue
        line = _BULLET.sub("", line)
        line = _NUMBERED.sub("", line)
        line = line.strip()
        if not line:
            continue
        if current is None:
            top.append(line)
        else:
            sections[current].append(line)
    return top, sections


def _extract_name_and_contact(top: list[str]) -> dict[str, str | None]:
    draft: dict[str, str | None] = {
        "full_name": None,
        "location": None,
        "phone": None,
        "email": None,
        "linkedin_url": None,
        "portfolio_url": None,
    }

    for line in top:
        low = line.lower().strip(" .")
        if low in _DOCUMENT_HEADERS or "@" in line or re.search(r"\d", line):
            continue
        if _NAME_RE.match(line) or _ALL_CAPS_NAME_RE.match(line):
            if len(line.split()) <= 6:
                draft["full_name"] = line
                break

    for line in top:
        stripped = line.strip().strip(",")
        if "," in stripped and "@" not in stripped and not re.search(r"\d", stripped):
            if len(stripped) <= 80:
                draft["location"] = stripped
                break

    full = "\n".join(top)
    email = _EMAIL_RE.search(full)
    if email:
        draft["email"] = email.group(0)
    phone = _PHONE_RE.search(full)
    if phone:
        digits = sum(ch.isdigit() for ch in phone.group(0))
        if 7 <= digits <= 16:
            cleaned = phone.group(0).strip()
            if len(cleaned) <= 24:
                draft["phone"] = cleaned
    linkedin = _LINKEDIN_RE.search(full)
    if linkedin:
        link = linkedin.group(0)
        if not link.lower().startswith("http"):
            link = "https://" + link
        draft["linkedin_url"] = link
    github = _GITHUB_RE.search(full)
    if github and not draft["linkedin_url"]:
        draft["portfolio_url"] = github.group(0)
    return draft


def _split_items(lines: list[str], max_items: int) -> list[str]:
    items: list[str] = []
    seen: set[str] = set()
    for line in lines:
        if not line:
            continue
        for part in re.split(r"[,|·•]|\t", line):
            part = part.strip().strip("*•·-–—")
            if not part or len(part) > 100:
                continue
            key = part.lower()
            if key in seen:
                continue
            seen.add(key)
            items.append(part)
            if len(items) >= max_items:
                return items
    return items


def _parse_summary(lines: list[str]) -> str | None:
    text = " ".join(part for part in lines if part).strip()
    text = re.sub(r"\s+", " ", text)
    if len(text) < 12:
        return None
    return text[:4000]


def _date_range_for(line: str) -> tuple[re.Match[str], tuple[str, str]] | None:
    """Return (match, (start, end)) for a date range found in a line."""
    m = _RANGE_RE.search(line)
    if not m:
        return None
    start = f"{m.group(1) or ''}{m.group(2)}".strip()
    end = f"{m.group(3) or ''}".strip()
    return m, (start, end)


def _split_paragraphs(lines: list[str]) -> list[list[str]]:
    paragraphs: list[list[str]] = []
    current: list[str] = []
    for line in lines:
        if not line:
            if current:
                paragraphs.append(current)
                current = []
            continue
        current.append(line)
    if current:
        paragraphs.append(current)
    return [p for p in paragraphs if any(l for l in p if l.strip())]


def _parse_experience(lines: list[str]) -> list[dict[str, Any]]:
    entries: list[dict[str, Any]] = []
    for para in _split_paragraphs(lines):
        entry: dict[str, Any] = {
            "employer": None,
            "title": None,
            "start_date": None,
            "end_date": None,
            "description": None,
            "source": SOURCE,
        }
        found = None
        for i, line in enumerate(para):
            res = _date_range_for(line)
            if res is not None:
                found = (i, res)
                break

        if found is None:
            # No reliable date structure — keep the paragraph as a single
            # reviewable text block without inventing company/title/dates.
            desc = " ".join(l for l in para if l.strip()).strip()
            if desc:
                entry["description"] = desc[:2000]
            entries.append(entry)
            continue

        idx, (m, (start, end)) = found
        entry["start_date"] = start[:50]
        entry["end_date"] = end[:50]

        # The date line usually also carries the employer before the range.
        leftover = para[idx][: m.start()].strip()
        leftover = leftover.rstrip(":|/·•-–—").rstrip("|").strip()
        title_idx: int | None = None
        if leftover:
            entry["employer"] = leftover[:255]
            for j in range(idx + 1, len(para)):
                if not re.search(r"\d", para[j]) and len(para[j]) <= 120:
                    entry["title"] = para[j][:255]
                    title_idx = j
                    break
            if entry["title"] == entry["employer"]:
                entry["title"] = None
                title_idx = None

        desc_parts: list[str] = []
        for j, line in enumerate(para):
            if j == idx:
                continue
            if j == title_idx:
                continue
            head = _BULLET.sub("", line).strip()
            if head == entry["title"]:
                continue
            if head:
                desc_parts.append(head)
        desc = " ".join(desc_parts).strip()
        if desc:
            entry["description"] = desc[:2000]
        if not entry["employer"] and not entry["title"] and not entry["description"]:
            continue
        entries.append(entry)
        if len(entries) >= 50:
            break
    return entries


def _parse_education(lines: list[str]) -> list[dict[str, Any]]:
    entries: list[dict[str, Any]] = []
    for para in _split_paragraphs(lines):
        segments: list[str] = []
        for line in para:
            for seg in re.split(r"[|/•]", line):
                seg = seg.strip().strip("•·-–—")
                if seg:
                    segments.append(seg)

        start = end = None
        remaining: list[str] = []
        for seg in segments:
            m = _RANGE_RE.search(seg)
            if m and start is None:
                start = f"{m.group(1) or ''}{m.group(2)}".strip()
                end = m.group(3).strip()
                seg = (seg[: m.start()] + seg[m.end():]).strip().strip(",;:|-–—")
                if seg:
                    remaining.append(seg)
                continue
            remaining.append(seg)

        degree = None
        institution = None
        for seg in remaining:
            if not seg:
                continue
            if re.search(r"\d", seg):
                continue
            lowered = seg.lower()
            if any(kw in lowered for kw in _DEGREE_KEYWORDS if len(kw) > 2) or (
                re.search(r"\b(bs|ba|bsc|ms|ma|mba|phd)\b", lowered)
            ):
                degree = degree or seg
                continue
            if institution is None and len(seg) <= 120:
                institution = seg

        entry: dict[str, Any] = {
            "institution": None,
            "degree": None,
            "field": None,
            "start_date": start,
            "end_date": end,
            "source": SOURCE,
        }
        if degree:
            entry["degree"] = degree[:255]
        if institution and institution != degree:
            entry["institution"] = institution[:255]
        if not entry["institution"] and not entry["degree"]:
            continue  # ambiguous — leave unresolved rather than guess
        entries.append(entry)
        if len(entries) >= 20:
            break
    return entries


def _parse_certifications(lines: list[str]) -> list[dict[str, Any]]:
    certs: list[dict[str, Any]] = []
    for line in lines:
        if not line:
            continue
        text = line.strip()
        year = None
        m = re.search(r"(?:[-–—]\s*|\()\s*((?:19|20)\d{2})\s*\)?\s*$", text)
        if m:
            year = m.group(1)
            text = text[: m.start()].strip().rstrip("(-–—, ").strip()
        issuer = None
        m2 = re.match(r"^(.+?)\s*\(\s*([^()]+?)\s*\)\s*$", text)
        if m2 and " " in m2.group(2):
            text = m2.group(1).strip()
            issuer = m2.group(2).strip()
        elif m2:
            text = m2.group(1).strip()
            issuer = m2.group(2).strip()
        if not text or len(text) > 200:
            continue
        item: dict[str, Any] = {"name": text, "issuer": None, "year": None, "source": SOURCE}
        if issuer and len(issuer) <= 200:
            item["issuer"] = issuer
        if year:
            item["year"] = year
        certs.append(item)
        if len(certs) >= 30:
            break
    return certs


def _parse_languages(lines: list[str]) -> list[dict[str, Any]]:
    langs: list[dict[str, Any]] = []
    for raw in _split_items(lines, max_items=30):
        item: dict[str, Any] = {"name": raw, "level": None, "source": SOURCE}
        m = re.match(r"^(.+?)\s*\(\s*([^()]{1,40})\s*\)\s*$", raw)
        if m:
            item["name"] = m.group(1).strip()
            item["level"] = m.group(2).strip()
        langs.append(item)
    return langs


def _parse_text(text: str) -> dict[str, Any]:
    top, sections = _segment_sections(text)
    draft = _empty_draft()
    draft.update(_extract_name_and_contact(top))

    if sections.get("summary"):
        draft["professional_summary"] = _parse_summary(sections["summary"])
    if sections.get("skills"):
        draft["skills"] = [
            {"name": name, "source": SOURCE}
            for name in _split_items(sections["skills"], max_items=200)
        ]
    if sections.get("experience"):
        draft["experience"] = _parse_experience(sections["experience"])
    if sections.get("education"):
        draft["education"] = _parse_education(sections["education"])
    if sections.get("certifications"):
        draft["certifications"] = _parse_certifications(sections["certifications"])
    if sections.get("languages"):
        draft["languages"] = _parse_languages(sections["languages"])
    # Projects are intentionally not auto-extracted: the user adds them.
    draft["projects"] = []
    return draft